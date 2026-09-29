from pathlib import Path
from io import BytesIO
import json
import re

import numpy as np
import pandas as pd
import streamlit as st
from docx import Document
from pypdf import PdfReader
from rank_bm25 import BM25Okapi
from sentence_transformers import SentenceTransformer


st.set_page_config(
    page_title="Explainable Candidate–Job Matching",
    layout="wide",
)

ROOT = Path(__file__).resolve().parent
ONTOLOGY_PATH = ROOT / "skill_ontology.json"


@st.cache_resource
def load_model():
    return SentenceTransformer("all-MiniLM-L6-v2")


@st.cache_data
def load_ontology():
    with ONTOLOGY_PATH.open("r", encoding="utf-8") as handle:
        ontology = json.load(handle)

    if not isinstance(ontology, dict) or not ontology:
        raise ValueError("skill_ontology.json must contain a skill dictionary.")

    return ontology


def read_document(uploaded_file):
    """Extract machine-readable text from TXT, PDF or DOCX."""
    content = uploaded_file.getvalue()
    extension = Path(uploaded_file.name).suffix.lower()

    if extension == ".txt":
        return content.decode("utf-8", errors="replace")

    if extension == ".pdf":
        reader = PdfReader(BytesIO(content))
        return "\n".join(page.extract_text() or "" for page in reader.pages)

    if extension == ".docx":
        document = Document(BytesIO(content))
        paragraphs = [paragraph.text for paragraph in document.paragraphs]
        for table in document.tables:
            for row in table.rows:
                paragraphs.append(" ".join(cell.text for cell in row.cells))
        return "\n".join(paragraphs)

    raise ValueError(f"Unsupported file type: {extension}")


def clean_text(text):
    text = text.replace("\u00ad", "")
    text = re.sub(r"-\s*\n\s*(?=\w)", "", text)
    return re.sub(r"\s+", " ", text).strip()


def remove_direct_identifiers(text):
    """Basic minimisation; this does not guarantee anonymisation."""
    text = re.sub(
        r"\b[A-Z0-9._%+-]+@[A-Z0-9.-]+\.[A-Z]{2,}\b",
        "[email removed]",
        text,
        flags=re.IGNORECASE,
    )
    text = re.sub(
        r"(?<!\w)(?:\+?\d[\d\s().-]{8,}\d)(?!\w)",
        "[phone removed]",
        text,
    )
    return text


def alias_pattern(alias):
    """Match complete terms while allowing flexible spaces."""
    parts = re.split(r"\s+", alias.strip())
    body = r"\s+".join(re.escape(part) for part in parts)
    return re.compile(
        rf"(?<![A-Za-z0-9]){body}(?![A-Za-z0-9])",
        flags=re.IGNORECASE,
    )


def is_negated(text, match_start):
    preceding = text[max(0, match_start - 80):match_start].lower()
    preceding = re.split(r"[.;\n]", preceding)[-1]
    return bool(
        re.search(
            r"(?:no experience (?:with|in|of)|"
            r"no knowledge (?:of|in)|"
            r"not experienced (?:with|in)|"
            r"without experience (?:with|in)|"
            r"lack(?:s|ing)? experience (?:with|in)|"
            r"unfamiliar with)\s*$",
            preceding,
        )
    )


def extract_skills(text, ontology):
    found = set()
    evidence = {}

    for canonical, aliases in ontology.items():
        terms = {str(canonical), *map(str, aliases)}

        for term in sorted(terms, key=len, reverse=True):
            for match in alias_pattern(term).finditer(text):
                if is_negated(text, match.start()):
                    continue

                found.add(canonical)
                evidence.setdefault(canonical, []).append(
                    text[max(0, match.start() - 45):
                         min(len(text), match.end() + 65)].strip()
                )

            if canonical in found:
                break

    return found, evidence


def job_skill_sections(text):
    """Identify required and preferred sections when the headings are present."""
    required_heading = re.search(
        r"\b(required|essential|must[- ]have|minimum)\b"
        r"(?:\s+(?:skills|requirements|qualifications))?\s*:",
        text,
        flags=re.IGNORECASE,
    )
    preferred_heading = re.search(
        r"\b(preferred|desirable|nice[- ]to[- ]have)\b"
        r"(?:\s+(?:skills|requirements|qualifications))?\s*:",
        text,
        flags=re.IGNORECASE,
    )

    if required_heading and preferred_heading:
        if required_heading.start() < preferred_heading.start():
            required_text = text[
                required_heading.end():preferred_heading.start()
            ]
            preferred_text = text[preferred_heading.end():]
        else:
            preferred_text = text[
                preferred_heading.end():required_heading.start()
            ]
            required_text = text[required_heading.end():]
        return required_text, preferred_text

    # If headings are absent, all identified skills are treated as required.
    return text, ""


def tokenize(text):
    return re.findall(r"[a-z0-9]+", text.lower())


st.title("Explainable Candidate–Job Matching")
st.write(
    "Upload one job description and at least two CVs. The prototype "
    "calculates Sentence-BERT similarity, structured skill coverage, "
    "BM25 scores and a 60:40 hybrid ranking."
)
st.caption(
    "Research prototype only. Review the original CVs and job criteria "
    "before making any decision."
)

if not ONTOLOGY_PATH.exists():
    st.error(
        "Upload skill_ontology.json to the same GitHub folder as app.py."
    )
    st.stop()

try:
    ontology = load_ontology()
except Exception as exc:
    st.error(f"Could not load the skill dictionary: {exc}")
    st.stop()

job_file = st.file_uploader(
    "Upload job description",
    type=["txt", "pdf", "docx"],
)
cv_files = st.file_uploader(
    "Upload two or more CVs",
    type=["txt", "pdf", "docx"],
    accept_multiple_files=True,
)

if st.button("Rank candidates", type="primary"):
    if job_file is None or len(cv_files) < 2:
        st.warning("Upload one job description and at least two CVs.")
        st.stop()

    try:
        job_text = clean_text(read_document(job_file))
        cv_texts = [
            clean_text(remove_direct_identifiers(read_document(file)))
            for file in cv_files
        ]
    except Exception as exc:
        st.error(f"Could not read a document: {exc}")
        st.stop()

    if len(job_text) < 30:
        st.error(
            "The job description contains too little extractable text. "
            "Scanned PDFs need OCR or a machine-readable copy."
        )
        st.stop()

    unreadable = [
        file.name
        for file, text in zip(cv_files, cv_texts)
        if len(text) < 30
    ]
    if unreadable:
        st.error(
            "These CVs contain too little extractable text: "
            + ", ".join(unreadable)
        )
        st.stop()

    required_text, preferred_text = job_skill_sections(job_text)
    required_skills, _ = extract_skills(required_text, ontology)
    preferred_skills, _ = extract_skills(preferred_text, ontology)
    preferred_skills -= required_skills

    if not required_skills:
        st.warning(
            "No skills from the fixed dictionary were identified in the "
            "job description. The skill component cannot be calculated "
            "meaningfully for this job."
        )
        st.stop()

    with st.spinner("Loading Sentence-BERT and scoring documents..."):
        model = load_model()
        embeddings = model.encode(
            [job_text] + cv_texts,
            normalize_embeddings=True,
            convert_to_numpy=True,
        )
        cosine_scores = embeddings[1:] @ embeddings[0]
        semantic_scores = np.clip(
            (cosine_scores + 1.0) / 2.0, 0.0, 1.0
        )

        corpus = [tokenize(text) for text in cv_texts]
        bm25 = BM25Okapi(corpus)
        bm25_scores = bm25.get_scores(tokenize(job_text))

    rows = []
    evidence_by_candidate = {}

    for index, (file, cv_text) in enumerate(zip(cv_files, cv_texts)):
        candidate_skills, evidence = extract_skills(cv_text, ontology)

        matched_required = sorted(required_skills & candidate_skills)
        missing_required = sorted(required_skills - candidate_skills)
        matched_preferred = sorted(preferred_skills & candidate_skills)

        required_coverage = (
            len(matched_required) / len(required_skills)
        )
        preferred_coverage = (
            len(matched_preferred) / len(preferred_skills)
            if preferred_skills else 0.0
        )

        skill_score = (
            0.80 * required_coverage
            + 0.20 * preferred_coverage
        )
        hybrid_score = (
            0.60 * float(semantic_scores[index])
            + 0.40 * skill_score
        )

        candidate_id = f"CV{index + 1:03d}"
        evidence_by_candidate[candidate_id] = {
            "file": file.name,
            "matched_required": matched_required,
            "missing_required": missing_required,
            "matched_preferred": matched_preferred,
            "evidence": evidence,
        }

        rows.append({
            "candidate_id": candidate_id,
            "cv_file": file.name,
            "semantic_score": float(semantic_scores[index]),
            "required_coverage": required_coverage,
            "preferred_coverage": preferred_coverage,
            "skill_score": skill_score,
            "bm25_score": float(bm25_scores[index]),
            "hybrid_score": hybrid_score,
        })

    result = pd.DataFrame(rows)
    result = result.sort_values(
        ["hybrid_score", "candidate_id"],
        ascending=[False, True],
        kind="stable",
    ).reset_index(drop=True)
    result.insert(0, "rank", range(1, len(result) + 1))

    st.session_state["result"] = result
    st.session_state["evidence"] = evidence_by_candidate
    st.session_state["job_skills"] = {
        "required": sorted(required_skills),
        "preferred": sorted(preferred_skills),
    }

if "result" in st.session_state:
    result = st.session_state["result"]
    evidence_by_candidate = st.session_state["evidence"]
    job_skills = st.session_state["job_skills"]

    st.subheader("New candidate ranking")
    st.dataframe(
        result.round(4),
        use_container_width=True,
        hide_index=True,
    )

    st.download_button(
        "Download ranking CSV",
        data=result.to_csv(index=False).encode("utf-8"),
        file_name="candidate_ranking.csv",
        mime="text/csv",
    )

    with st.expander("Skills identified in the job description"):
        st.write("**Required:**", job_skills["required"])
        st.write("**Preferred:**", job_skills["preferred"])

    selected_id = st.selectbox(
        "Inspect a candidate",
        result["candidate_id"].tolist(),
    )
    details = evidence_by_candidate[selected_id]

    st.write(f"**File:** {details['file']}")
    left, right = st.columns(2)

    with left:
        st.write(
            "**Matched required skills:**",
            details["matched_required"] or "None identified",
        )
        st.write(
            "**Matched preferred skills:**",
            details["matched_preferred"] or "None identified",
        )

    with right:
        st.write(
            "**Missing required skills:**",
            details["missing_required"] or "None identified",
        )

    with st.expander("Text evidence for matched skills"):
        for skill in (
            details["matched_required"]
            + details["matched_preferred"]
        ):
            st.write(f"**{skill}**")
            for snippet in details["evidence"].get(skill, [])[:2]:
                st.write(f"• {snippet}")

    st.info(
        "Scores are matching aids, not suitability decisions. "
        "The skill dictionary can miss synonyms or context, and "
        "Sentence-BERT similarity does not prove competence. "
        "The BM25 score is displayed separately and is not part "
        "of the hybrid score."
    )
